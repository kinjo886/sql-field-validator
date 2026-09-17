import re
import os
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinterdnd2 import TkinterDnD, DND_FILES
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# =================全局常量================
EXCLUDE_TABLE_NAME = "tab_name"
REG_TABLE_COMMENT = re.compile(r'--\s*表：\s*([a-zA-Z0-9_]+)\.([a-zA-Z0-9_]+)', re.IGNORECASE)
# 严格字段正则边界，匹配独立字段
def get_strict_field_regex(field):
    return re.compile(rf'(?<!\w){re.escape(field)}(?!\w)', re.IGNORECASE)

def parse_sql_content(full_text, file_full_path):
    result_list = []
    file_short_name = os.path.basename(file_full_path)
    matches = list(REG_TABLE_COMMENT.finditer(full_text))
    skip_example_count = 0

    for idx, match in enumerate(matches):
        db_name = match.group(1)
        tbl_name = match.group(2)

        if tbl_name.strip() == EXCLUDE_TABLE_NAME:
            skip_example_count += 1
            continue

        start_pos = match.start()
        if idx + 1 < len(matches):
            end_pos = matches[idx+1].start()
        else:
            end_pos = len(full_text)

        table_sql_block = full_text[start_pos:end_pos]
        if "create table" not in table_sql_block.lower():
            skip_example_count += 1
            continue

        result_list.append({
            "source_file": file_short_name,
            "database": db_name,
            "table": tbl_name,
            "sql_block": table_sql_block
        })
    return result_list

def export_excel(all_data, target_column, save_path, strict_match=False, only_missing=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "检查表结果"
    headers = [
        "序号",
        "来源SQL文件",
        "数据库名",
        "表名",
        f"是否存在【{target_column}】",
        "匹配上下文片段"
    ]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    for col in range(1, len(headers)+1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    target_low = target_column.lower()
    strict_re = get_strict_field_regex(target_column) if strict_match else None

    seq_out = 1
    for row_data in all_data:
        block = row_data["sql_block"]
        if strict_re:
            has_field = True if strict_re.search(block) else False
        else:
            has_field = target_low in block.lower()

        # 如果勾选【仅导出缺失】，跳过存在字段的表
        if only_missing and has_field:
            continue

        flag_text = "✅ 存在" if has_field else "❌ 不存在"
        context_snippet = ""
        if has_field:
            pos = block.lower().find(target_low)
            s = max(0, pos - 180)
            e = min(len(block), pos + len(target_low)+180)
            context_snippet = block[s:e].replace("\n", " ").strip()

        row = [
            seq_out,
            row_data["source_file"],
            row_data["database"],
            row_data["table"],
            flag_text,
            context_snippet
        ]
        ws.append(row)
        status_cell = ws.cell(row=seq_out+1, column=5)
        if has_field:
            status_cell.font = Font(color="009933", bold=True)
        else:
            status_cell.font = Font(color="CC0000")
        seq_out += 1

    widths = [8, 32, 22, 24, 22, 95]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(save_path)
    return seq_out - 1

# =================GUI界面逻辑================
class App(TkinterDnD.Tk):
    def __init__(self):
        super().__init__()
        self.title("SQL建表字段核查工具")
        self.geometry("760x520")
        self.sql_file_list = []
        self.custom_output_dir = None
        self.is_running = False # 是否正在执行任务

        main_frame = ttk.Frame(self, padding=15)
        main_frame.pack(fill=tk.BOTH, expand=True)

        # 第1行 字段名称
        ttk.Label(main_frame, text="待查询字段名称：").grid(row=0, column=0, sticky="w")
        self.entry_field = ttk.Entry(main_frame, width=48)
        self.entry_field.grid(row=0, column=1, padx=10, pady=5)
        self.entry_field.insert(0, "<INTERNAL_DATASET>")

        # 第2行 输出目录
        ttk.Label(main_frame, text="Excel输出目录：").grid(row=1, column=0, sticky="w", pady=(6,0))
        self.var_output_path = tk.StringVar(value="【默认】Excel生成在第一个SQL文件同目录")
        ttk.Label(main_frame, textvariable=self.var_output_path, foreground="#204080").grid(row=1, column=1, padx=10, sticky="we")
        btn_output = ttk.Button(main_frame, text="选择目录", command=self.select_output_dir)
        btn_output.grid(row=1, column=2, padx=(5,0))
        btn_reset_path = ttk.Button(main_frame, text="恢复默认", command=self.reset_output_dir)
        btn_reset_path.grid(row=1, column=3, padx=(3,0))

        # 第3行 检索选项
        opt_frame = ttk.Frame(main_frame)
        opt_frame.grid(row=2, column=0, columnspan=4, sticky="w", pady=(4,0))
        ttk.Label(opt_frame, text="匹配模式：").pack(side=tk.LEFT)
        self.match_mode = tk.StringVar()
        cbx_mode = ttk.Combobox(opt_frame, textvariable=self.match_mode, width=18, state="readonly")
        cbx_mode["values"] = ["简单包含匹配", "严格独立字段匹配"]
        cbx_mode.current(0)
        self.mode_map = {
            "简单包含匹配":"simple",
            "严格独立字段匹配":"strict"
        }
        cbx_mode.pack(side=tk.LEFT, padx=(5,15))

        self.var_only_miss = tk.BooleanVar()
        chk_miss = ttk.Checkbutton(opt_frame, text="仅导出【不包含字段】的表", variable=self.var_only_miss)
        chk_miss.pack(side=tk.LEFT)

        # 第4行 拖拽提示
        ttk.Label(main_frame, text="👇下方区域拖拽多个.sql文件，或点击添加文件").grid(row=3, column=0, columnspan=4, sticky="w", pady=(10,0))
        self.listbox = tk.Listbox(main_frame, width=90, height=11, selectmode=tk.EXTENDED)
        self.listbox.grid(row=4, column=0, columnspan=4, pady=5)
        self.listbox.drop_target_register(DND_FILES)
        self.listbox.dnd_bind('<<Drop>>', self.drop_files)

        # 第5行按钮组
        btn_frame = ttk.Frame(main_frame)
        btn_frame.grid(row=5, column=0, columnspan=4, pady=8)
        self.btn_add = ttk.Button(btn_frame, text="添加SQL文件", command=self.select_files)
        self.btn_add.pack(side=tk.LEFT, padx=5)
        self.btn_del = ttk.Button(btn_frame, text="删除选中", command=self.delete_selected)
        self.btn_del.pack(side=tk.LEFT, padx=5)
        self.btn_clear = ttk.Button(btn_frame, text="清空列表", command=self.clear_list)
        self.btn_clear.pack(side=tk.LEFT, padx=5)
        self.btn_run = ttk.Button(btn_frame, text="开始分析并导出Excel", command=self.run_task)
        self.btn_run.pack(side=tk.LEFT, padx=15)

        # 日志
        ttk.Label(main_frame, text="运行日志：").grid(row=6, column=0, sticky="w")
        self.log_text = tk.Text(main_frame, height=5, width=92)
        self.log_text.grid(row=7, column=0, columnspan=4)

    def log(self, msg):
        self.log_text.insert(tk.END, msg + "\n")
        self.log_text.see(tk.END)
        self.update_idletasks()

    def select_output_dir(self):
        folder = filedialog.askdirectory(title="选择Excel结果保存文件夹")
        if folder:
            self.custom_output_dir = folder
            self.var_output_path.set(folder)

    def reset_output_dir(self):
        self.custom_output_dir = None
        self.var_output_path.set("【默认】Excel生成在第一个SQL文件同目录")

    def delete_selected(self):
        sel = self.listbox.curselection()
        if not sel:
            messagebox.showinfo("提示", "请先在列表选中要删除的文件！")
            return
        # 倒序删除防止索引错乱
        for idx in reversed(sel):
            fname_display = self.listbox.get(idx)
            remove_idx = None
            for i, fullpath in enumerate(self.sql_file_list):
                if os.path.basename(fullpath) == fname_display:
                    remove_idx = i
                    break
            if remove_idx is not None:
                del self.sql_file_list[remove_idx]
            self.listbox.delete(idx)
        self.log("✅ 已删除选中条目")

    def drop_files(self, event):
        paths = self.tk.splitlist(event.data)
        for p in paths:
            if p.lower().endswith(".sql") and p not in self.sql_file_list:
                self.sql_file_list.append(p)
                self.listbox.insert(tk.END, os.path.basename(p))
        self.log(f"✅ 成功拖入 {len(paths)} 个文件")

    def select_files(self):
        paths = filedialog.askopenfilenames(filetypes=[("SQL脚本", "*.sql"), ("所有文件", "*.*")])
        for p in paths:
            if p not in self.sql_file_list:
                self.sql_file_list.append(p)
                self.listbox.insert(tk.END, os.path.basename(p))
        self.log(f"✅ 选择文件，当前共{len(self.sql_file_list)}个SQL")

    def clear_list(self):
        self.sql_file_list.clear()
        self.listbox.delete(0, tk.END)
        self.log_text.delete(1.0, tk.END)

    def set_widget_state(self, enable: bool):
        """任务执行时锁定/解锁按钮"""
        state = tk.NORMAL if enable else tk.DISABLED
        self.btn_add.config(state=state)
        self.btn_del.config(state=state)
        self.btn_clear.config(state=state)
        self.btn_run.config(state=state)
        self.is_running = not enable

    def run_task(self):
        target_col = self.entry_field.get().strip()
        if not target_col:
            messagebox.showwarning("提示", "请输入要查询的字段名！")
            return
        if len(self.sql_file_list) == 0:
            messagebox.showwarning("提示", "请拖拽/选择至少一个SQL文件！")
            return
        if self.is_running:
            messagebox.showinfo("提示", "任务正在执行，请稍等！")
            return

        select_text = self.match_mode.get()
        strict_mode = True if self.mode_map[select_text] == "strict" else False
        only_missing = self.var_only_miss.get()
        self.set_widget_state(False)
        all_result = []
        self.log("====开始解析SQL文件====")
        try:
            for file_path in self.sql_file_list:
                fname = os.path.basename(file_path)
                try:
                    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                        content = f.read()
                except Exception as e:
                    self.log(f"❌【{fname}】读取失败：{str(e)}")
                    continue
                table_data = parse_sql_content(content, file_path)
                all_result.extend(table_data)
                self.log(f"✅【{fname}】解析完成，有效表数量：{len(table_data)}")

            if len(all_result) == 0:
                self.log("未识别到任何有效建表信息！")
                messagebox.showerror("结束", "没有读取到有效数据表！")
                return

            excel_filename = f"SQL字段核查结果_{target_col}.xlsx"
            if self.custom_output_dir is not None:
                out_dir = self.custom_output_dir
            else:
                out_dir = os.path.dirname(self.sql_file_list[0])
            out_file = os.path.join(out_dir, excel_filename)

            export_row_count = export_excel(all_result, target_col, out_file, strict_mode, only_missing)

            # 统计命中数量
            hit = 0
            strict_re = get_strict_field_regex(target_col)
            target_low = target_col.lower()
            for x in all_result:
                if strict_mode:
                    if strict_re.search(x["sql_block"]):
                        hit +=1
                else:
                    if target_low in x["sql_block"].lower():
                        hit +=1

            total = len(all_result)
            self.log(f"\n====分析完成====")
            self.log(f"总共有效表：{total} 张")
            self.log(f"包含字段【{target_col}】：{hit} 张")
            self.log(f"不包含字段：{total-hit} 张")
            self.log(f"Excel本次输出条目：{export_row_count} 条")
            self.log(f"📁 文件输出位置：{out_file}")
            messagebox.showinfo("执行完成", f"分析完毕！\n总共{total}张表\n包含字段：{hit}张\nExcel已生成:\n{out_file}")
        finally:
            self.set_widget_state(True)

if __name__ == "__main__":
    app = App()
    app.mainloop()