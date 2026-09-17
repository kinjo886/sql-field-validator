import re
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# =====================【配置区 务必修改】=====================
# 你的两个SQL文件完整路径
SQL_FILES = [
    r"D:\Desktop\公司\数据验证\all_target_db_ddl.sql",
    r"D:\Desktop\公司\数据验证\all_target_db_ddl1.sql"
]
# 需要查找的字段名
TARGET_COLUMN = "<INTERNAL_DATASET>"
# Excel输出路径
OUTPUT_EXCEL = r"D:\Desktop\公司\数据验证\表字段核查结果.xlsx"
# 【黑名单】要排除的示例占位表名
EXCLUDE_TABLE_NAME = "tab_name"
# ============================================================

# 正则匹配注释： -- 表： dbname.tabname
REG_TABLE_COMMENT = re.compile(r'--\s*表：\s*([a-zA-Z0-9_]+)\.([a-zA-Z0-9_]+)', re.IGNORECASE)

def parse_sql_content(full_text, file_full_path):
    result_list = []
    # 获取单纯文件名（不带路径）
    file_short_name = os.path.basename(file_full_path)
    matches = list(REG_TABLE_COMMENT.finditer(full_text))
    total_tables = len(matches)
    skip_example_count = 0
    print(f"[{file_short_name}] 原始识别注释标记数量：{total_tables}")

    for idx, match in enumerate(matches):
        db_name = match.group(1)
        tbl_name = match.group(2)

        # ==========过滤示例占位tab_name==========
        if tbl_name.strip() == EXCLUDE_TABLE_NAME:
            skip_example_count += 1
            continue

        start_pos = match.start()
        if idx + 1 < len(matches):
            end_pos = matches[idx+1].start()
        else:
            end_pos = len(full_text)

        table_sql_block = full_text[start_pos:end_pos]

        # 区间内不存在CREATE TABLE直接跳过
        if "create table" not in table_sql_block.lower():
            skip_example_count += 1
            continue

        target_low = TARGET_COLUMN.lower()
        block_low = table_sql_block.lower()
        has_field = target_low in block_low

        context_snippet = ""
        if has_field:
            pos = block_low.find(target_low)
            s = max(0, pos - 180)
            e = min(len(table_sql_block), pos + len(target_low)+180)
            context_snippet = table_sql_block[s:e].replace("\n", " ").strip()

        result_list.append({
            "source_file": file_short_name,   # 这里只保存文件名！
            "database": db_name,
            "table": tbl_name,
            "has_target": has_field,
            "sql_context": context_snippet
        })
    print(f"[{file_short_name}] 过滤掉示例占位表数量：{skip_example_count}，有效真实表：{len(result_list)}")
    return result_list


def main():
    all_data = []
    for file_path in SQL_FILES:
        print(f"\n正在读取文件：{file_path}")
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
        except Exception as err:
            print(f"读取失败 {file_path} >>> {err}")
            continue
        table_info = parse_sql_content(content, file_path)
        all_data.extend(table_info)

    total = len(all_data)
    hit_count = sum(1 for x in all_data if x["has_target"])
    print(f"\n=====汇总结果=====")
    print(f"有效真实表总数：{total}")
    print(f"包含字段【{TARGET_COLUMN}】的表：{hit_count}")
    print(f"不包含字段的表：{total - hit_count}")

    # ==========写入Excel==========
    wb = Workbook()
    ws = wb.active
    ws.title = "检查表结果"
    headers = [
        "序号",
        "来源SQL文件",
        "数据库名",
        "表名",
        f"是否存在【{TARGET_COLUMN}】",
        "匹配上下文片段"
    ]
    ws.append(headers)

    # 表头样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    for col in range(1, len(headers)+1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 填充每行数据
    for seq, row_data in enumerate(all_data, start=1):
        flag_text = "✅ 存在" if row_data["has_target"] else "❌ 不存在"
        row = [
            seq,
            row_data["source_file"],
            row_data["database"],
            row_data["table"],
            flag_text,
            row_data["sql_context"]
        ]
        ws.append(row)
        # 状态文字上色
        status_cell = ws.cell(row=seq+1, column=5)
        if row_data["has_target"]:
            status_cell.font = Font(color="009933", bold=True)
        else:
            status_cell.font = Font(color="CC0000")

    # 设置列宽（来源文件列宽度缩小）
    widths = [8, 32, 22, 24, 22, 95]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(OUTPUT_EXCEL)
    print(f"\n✅ Excel文件已生成：{OUTPUT_EXCEL}")

if __name__ == "__main__":
    main()